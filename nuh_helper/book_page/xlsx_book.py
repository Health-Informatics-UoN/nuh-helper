import contextlib
import shutil
from collections.abc import Generator as gen
from pathlib import Path

import openpyxl
from openpyxl.cell.read_only import EmptyCell

from .book import BasePage, BaseWorkBook, row_val_t


class XLSXWorkBook(BaseWorkBook):
    def __init__(self, file: str | Path, read_only: bool = True) -> None:
        if isinstance(file, str):
            file = Path(file)
        if not file.is_file():
            raise RuntimeError(f"xlsx file not found {file}")
        self._file = file
        self._book = openpyxl.load_workbook(file, read_only=read_only, data_only=True)

    def keys(self) -> list[str]:
        return [page.title for page in self._book.worksheets]

    def _stream_copy_impl(self, into: str | Path) -> gen[tuple[str, int, row_val_t]]:
        if not isinstance(into, Path):
            into = Path(into)

        if not into.name.endswith(".xlsx"):
            into = into / self._file.name

        shutil.copy2(self._file, into)
        the_copy = openpyxl.load_workbook(into)

        for name in the_copy.sheetnames:
            page = the_copy[name]

            for r in range(page.max_row):
                data = [page.cell(r + 1, c + 1).value for c in range(page.max_column)]
                yield name, r, data
                for c, val in enumerate(data):
                    if val != page.cell(r + 1, c + 1).value:
                        page.cell(r + 1, c + 1).value = val
        the_copy.save(into)

    @contextlib.contextmanager
    def __getitem__(self, name: str) -> gen[BasePage]:

        if len(name) > 31:
            raise KeyError("excel limits you to 31 character page names")

        class XLSXPage(BasePage):
            def _stream_rows_impl(self) -> gen[tuple[int, list[str]]]:
                for row_idx, row_data in enumerate(self._page.iter_rows()):
                    yield (
                        row_idx,
                        [
                            None if isinstance(cell, EmptyCell) else cell.internal_value
                            for cell in row_data
                        ],
                    )

        yield XLSXPage(name, self._book[name])
