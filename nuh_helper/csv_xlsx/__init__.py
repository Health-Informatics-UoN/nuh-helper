import csv
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.worksheet import Worksheet


def csvs_into_xlsx(xlsx: Path, csvs: None | Path | list[Path] = None) -> None:
    """convert a group of csv files into a single excel document

    - xlsx: Path - required path where the xslx file should be created
    - csvs: None | Path | list[Path] - optional path to read the .csv files from. if
        it's a path, all csv files in that folder are used. if not given, then the
        output folder will be scanned for the files

    """

    if csvs is None:
        csvs = xlsx.parent

    if isinstance(csvs, Path):
        csvs = list(csvs.glob("*.csv"))

    assert isinstance(csvs, list)

    assert not xlsx.is_dir()

    workbook = Workbook()
    first_page = True

    for csv_file in csvs:
        assert csv_file.is_file()
        csv_name: str = csv_file.stem

        with csv_file.open() as file:
            data = list(
                csv.reader(
                    file.readlines(),
                    delimiter=("\t" if csv_file.name.endswith(".tsv") else ","),
                )
            )

        # write the data
        if first_page:
            page: Worksheet = workbook.active
            page.title = csv_name
            overwrite_page(page, data)
            first_page = False
        else:
            append_page(workbook, csv_name, data)

    assert not first_page
    workbook.save(xlsx)


def csvs_from_xlsx(xlsx: Path, csvs: None | Path = None) -> None:
    """convert an excel file into a group of csv files

    - xlsx: Path - required path to the xslx file
    - csvs: None | Path - optional path to store the .csv files. if absent they'll be
        stored in the same folder as the xlsx file.
    """
    if csvs is None:
        csvs = xlsx.parent

    book = load_workbook(xlsx)

    for name in book.sheetnames:
        with (csvs / f"{name}.csv").open("w") as file:
            data = csv.writer(file)
            page = book[name]
            data.writerows(
                [
                    [
                        page.cell(row + 1, col + 1).value
                        for col in range(page.max_column)
                    ]
                    for row in range(page.max_row)
                ]
            )


__all__ = [
    csvs_into_xlsx,
    csvs_from_xlsx,
]


def append_page(wb: Workbook, sheet_name: str, sheet_data: list[list]) -> None:
    ws = wb.create_sheet(sheet_name)
    overwrite_page(ws, sheet_data)


def overwrite_page(ws: Worksheet, sheet_data: list[list]) -> None:
    print(
        "TODO; clear rows if ws.max_row > 0 : ws.delete_rows(1, max_row) // then ws.append(row) for row in sheet_data"
    )
    for row_idx, row_data in enumerate(sheet_data, start=1):
        for col_idx, cell_data in enumerate(row_data, start=1):
            ws.cell(row=row_idx, column=col_idx, value=cell_data)
