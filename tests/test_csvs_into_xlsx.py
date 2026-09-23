import csv
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from nuh_helper.csv_xlsx import (
    append_page,
    csvs_from_xlsx,
    csvs_into_xlsx,
    overwrite_page,
)


class ExcelCellMismatch(Exception):
    def __init__(
        self, page_name: str, row: int, col: int, expected: any, obtained: any
    ) -> None:
        message = f"[{page_name=} {row}, {col}] {expected=} != {obtained=}"
        super().__init__(message)
        self._message = message
        self._page_name = page_name
        self._row = row
        self._col = col
        self._expected = expected
        self._obtained = obtained


def assert_xlsx_same(expected: Path, obtained: Path) -> None:
    expected_book = load_workbook(expected)
    obtained_book = load_workbook(obtained)

    assert expected_book.sheetnames == obtained_book.sheetnames

    for page_name in expected_book.sheetnames:
        expected_page = expected_book[page_name]
        obtained_page = obtained_book[page_name]

        assert expected_page.max_column == obtained_page.max_column
        assert expected_page.max_row == obtained_page.max_row

        for row in range(expected_page.max_row):
            row += 1

            for col in range(expected_page.max_column):
                col += 1

                expected_cell = str(expected_page.cell(row, col).value)
                obtained_cell = str(obtained_page.cell(row, col).value)

                if expected_cell != obtained_cell:
                    raise ExcelCellMismatch(
                        page_name, row, col, expected_cell, obtained_cell
                    )


def test_compare_xlsx(tmp_path: Path) -> None:
    """i'm going to need to compare xlsx files - this function does that"""

    # create two identical workbooks
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"

    overwrite_page(
        ws,
        [
            ["id", "name", "email", "city", "amount"],
            ["1", "Alice Smith", "alice@example.com", "London", "42.50"],
            ["2", "Bob Jones", "bob.jones@example.net", "Manchester", "17.99"],
            ["3", "Carla Ruiz", "carla.ruiz@example.org", "Madrid", "88.10"],
            ["4", "David Kim", "david.kim@example.com", "Seoul", "23.45"],
            ["5", "Ella Brown", "ella.brown@example.net", "Bristol", "56.78"],
            ["6", "Farid Khan", "farid.khan@example.org", "Dubai", "120.00"],
            ["7", "Grace Lee", "grace.lee@example.com", "Toronto", "9.99"],
        ],
    )

    append_page(
        wb,
        "foo-bar",
        [
            ["product", "quantity", "price"],
            ["Widget", "12", "4.99"],
            ["Gadget", "7", "15.50"],
            ["Thingamajig", "3", "22.00"],
            ["Doohickey", "18", "2.75"],
            ["Gizmo", "5", "9.99"],
            ["Contraption", "9", "12.30"],
            ["Apparatus", "1", "45.00"],
            ["Device", "14", "7.25"],
        ],
    )
    left = tmp_path / "left.xlsx"
    wb.save(left)

    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"

    overwrite_page(
        ws,
        [
            ["id", "name", "email", "city", "amount"],
            ["1", "Alice Smith", "alice@example.com", "London", "42.50"],
            ["2", "Bob Jones", "bob.jones@example.net", "Manchester", "17.99"],
            ["3", "Carla Ruiz", "carla.ruiz@example.org", "Madrid", "88.10"],
            ["4", "David Kim", "david.kim@example.com", "Seoul", "23.45"],
            ["5", "Ella Brown", "ella.brown@example.net", "Bristol", "56.78"],
            ["6", "Farid Khan", "farid.khan@example.org", "Dubai", "120.00"],
            ["7", "Grace Lee", "grace.lee@example.com", "Toronto", "9.99"],
        ],
    )

    append_page(
        wb,
        "foo-bar",
        [
            ["product", "quantity", "price"],
            ["Widget", "12", "4.99"],
            ["Gadget", "7", "15.50"],
            ["Thingamajig", "3", "22.00"],
            ["Doohickey", "18", "2.75"],
            ["Gizmo", "5", "9.99"],
            ["Contraption", "9", "12.30"],
            ["Apparatus", "1", "45.00"],
            ["Device", "14", "7.25"],
        ],
    )
    right = tmp_path / "right.xlsx"
    wb.save(right)

    # will pass
    assert_xlsx_same(left, right)

    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"

    overwrite_page(
        ws,
        [
            ["id", "name", "email", "city", "amount"],
            ["1", "Alice Smith", "alice@example.com", "London", "42.50"],
            ["2", "Bob Jones", "bob.jones@example.net", "Manchester", "17.99"],
            ["3", "Carla Ruiz", "carla.ruiz@example.org", "Madrid", "88.10"],
            ["4", "David Kim", "david.kim@example.com", "Seoul", "23.45"],
            ["5", "Ella Brown", "ella.brown@example.net", "Bristol", "56.78"],
            ["6", "Farid Khan", "farid.khan@example.org", "Dubai", "120.00"],
            ["7", "Grace Lee", "grace.lee@example.com", "Toronto", "9.99"],
        ],
    )

    append_page(
        wb,
        "foo-bar",
        [
            ["product", "quantity", "price"],
            ["Widget", "12", "4.99"],
            ["Gadget", "72e", "15.50"],
            ["Thingamajig", "3", "22.00"],
            ["Doohickey", "18", "2.75"],
            ["Gizmo", "5", "9.99"],
            ["Contraption", "9", "12.30"],
            ["Apparatus", "1", "45.00"],
            ["Device", "14", "7.25"],
        ],
    )
    third = tmp_path / "third.xlsx"
    wb.save(third)

    with pytest.raises(ExcelCellMismatch) as info:
        assert_xlsx_same(left, third)

    assert info.value._page_name == "foo-bar"
    assert info.value._row == 3
    assert info.value._col == 2
    assert info.value._expected == "7"
    assert info.value._obtained == "72e"


def test_combine(tmp_path: Path) -> None:
    data = Path(__file__).parent / "data/csvs_into_xlsx"

    result = tmp_path / "output.xlsx"

    csvs_into_xlsx(result, data)

    assert_xlsx_same(data / "combined.xlsx", result)


def test_extract(tmp_path: Path) -> None:
    data = Path(__file__).parent / "data/csvs_into_xlsx"

    csvs_from_xlsx(data / "combined.xlsx", tmp_path)

    # check if there is an expected file for each obtained one
    for obtained in tmp_path.glob("*.csv"):
        assert (data / obtained.name).is_file()

    # check it the other way and also check contents
    for expected_path in data.glob("*.csv"):
        obtained_path = tmp_path / expected_path.name
        assert obtained_path.is_file()
        with expected_path.open() as file:
            expected_data = list(csv.reader(file))
        with obtained_path.open() as file:
            obtained_data = list(csv.reader(file))

        assert expected_data == obtained_data


def test_tsv(tmp_path: Path) -> None:
    data = Path(__file__).parent / "data/csvs_into_xlsx"

    result = tmp_path / "output.xlsx"

    csvs_into_xlsx(result, sorted(data.glob("*.tsv")))

    print(f"output = ${result}")

    assert_xlsx_same(data / "combined.xlsx", result)
