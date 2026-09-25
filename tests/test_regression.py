from pathlib import Path

import openpyxl

from nuh_helper.date_shift import (
    shift_excel_dates_inplace,
)


def test_carrot_dates(tmp_path: Path) -> None:
    """test shifting a sheet from carrot"""

    source = Path(__file__).parent / "data/regression/carrot_dates.xlsx"
    target = tmp_path / source.name
    offsets_old = tmp_path / "user-offsets.old.csv"
    offsets_new = tmp_path / "user-offsets.new.csv"

    shift_excel_dates_inplace(
        input_file=str(source),
        output_file=str(target),
        patient_sheet="observation",
        patient_id_col="person_id",
        sheet_configs={
            "observation": {
                "patient_id_col": "person_id",
                "header_row": 0,
                "date_columns": ["observation_date", "observation_datetime"],
                "text_columns": [
                    "observation_id",
                    "observation_concept_id",
                    "observation_type_concept_id",
                    "qualifier_source_value",
                ],
            },
        },
        min_shift_days=-20,
        max_shift_days=-1,
        seed=14333,
        linking_table_path=offsets_old,
        linking_table_output=offsets_new,
    )

    print(f"{target=}")
    print(f"{target=}")
    print(f"{target=}")
    print(f"{target=}")
    print(f"{target=}")

    # Load the workbook / page
    book = openpyxl.load_workbook(target)
    assert "observation" in book.sheetnames
    page = book["observation"]

    expected = [
        [
            "observation_id",
            "person_id",
            "observation_concept_id",
            "observation_date",
            "observation_datetime",
            "observation_type_concept_id",
            "qualifier_source_value",
        ],
        ["1", "1", "27394", "1951-12-13 00:00:00", "1951-12-13 00:00:00", "0", None],
        ["2", "2", "25567", "1981-11-02 00:00:00", "1981-11-02 00:00:00", "0", None],
        ["3", "3", "26241", "1997-04-24 00:00:00", "1997-10-19 00:00:00", "0", None],
        ["4", "4", "25531", "1975-05-26 00:00:00", "1975-06-24 00:00:00", "0", None],
        ["5", "5", "27394", "1976-04-21 00:00:00", "1976-04-21 00:00:00", "0", None],
        ["6", "6", "25567", "1966-09-15 00:00:00", "1966-09-15 00:00:00", "0", None],
        ["7", "7", "25508", "1956-11-09 00:00:00", "1956-12-08 00:00:00", "0", None],
        ["8", "8", "27395", "1985-02-26 00:00:00", "1984-12-31 00:00:00", "0", None],
        ["9", "10", "27394", "1993-09-04 00:00:00", "1993-07-06 00:00:00", "0", None],
        ["10", "11", "27394", "1976-04-05 00:00:00", "1976-04-05 00:00:00", "0", None],
        ["11", "12", "26241", "1997-05-10 00:00:00", "1997-11-04 00:00:00", "0", None],
        ["12", "13", "25531", "1975-05-24 00:00:00", "1975-06-22 00:00:00", "0", None],
    ]

    expected_rows = iter(expected)

    for row_idx, row_val in enumerate(page.iter_rows()):
        obtained_row = [
            str(cell.value) if cell.value is not None else None for cell in row_val
        ]
        expected_row = next(expected_rows)

        assert expected_row == obtained_row, (
            f"mismatch at {row_idx}\n\t{expected_row=}\n\t{obtained_row=}"
        )
