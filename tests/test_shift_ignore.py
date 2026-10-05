from pathlib import Path

import pytest
from openpyxl import load_workbook

from nuh_helper import shift_excel_dates_inplace
from nuh_helper.date_shift import HiddenDate, PatientColumnMissing, ShiftFoundNonDate


def test_shift_exceptions(tmp_path: Path) -> None:
    """check that misconfigurations cause an error"""

    source_file = Path(__file__).parent / "data/shift_ignore/workbook.xlsx"
    output_path = tmp_path / "target.xlsx"
    linking_table_old = Path(__file__).parent / "data/shift_ignore/offsets.csv"
    linking_table_out = tmp_path / "linking_table_out.csv"

    sheet_configs = {
        "page-desc": "skip",
        "page-data": {
            "patient_id_col": "pid",
            "date_columns": ["dob"],
            "text_columns": ["top"],
            "header_row": 1,
            "skip_rows_after_header": [],
            "shift_exceptions": {"dob": ["mssing"]},
        },
    }
    with pytest.raises(RuntimeError) as info:
        shift_excel_dates_inplace(
            input_file=str(source_file),
            output_file=str(output_path),
            patient_sheet="page-data",
            sheet_configs=sheet_configs,
            min_shift_days=-20,
            max_shift_days=-1,
            seed=14333,
            linking_table_path=str(linking_table_old),
            linking_table_output=str(linking_table_out),
        )
    assert str(info.value) == (
        "shift_exceptions was removed, update sheet_name='page-data'"
    )


def test_ptid_in_date_columns(tmp_path: Path) -> None:
    """check that misconfigurations cause an error"""

    source_file = Path(__file__).parent / "data/shift_ignore/workbook.xlsx"
    output_path = tmp_path / "target.xlsx"
    linking_table_old = Path(__file__).parent / "data/shift_ignore/offsets.csv"
    linking_table_out = tmp_path / "linking_table_out.csv"

    sheet_configs = {
        "page-desc": "skip",
        "page-data": {
            "patient_id_col": "pid",
            "date_columns": ["dob", "pid"],
            "text_columns": ["top"],
            "header_row": 1,
            "skip_rows_after_header": [],
            "shift_ignore": {"dob": ["mssing"]},
        },
    }
    with pytest.raises(ValueError) as info:
        shift_excel_dates_inplace(
            input_file=str(source_file),
            output_file=str(output_path),
            patient_sheet="page-data",
            sheet_configs=sheet_configs,
            min_shift_days=-20,
            max_shift_days=-1,
            seed=14333,
            linking_table_path=str(linking_table_old),
            linking_table_output=str(linking_table_out),
        )
    assert str(info.value) == (
        "patient_id_col='pid' shouldn't be in date_columns of sheet_name='page-data'"
    )


def test_ptid_in_text_columns(tmp_path: Path) -> None:
    """check that misconfigurations cause an error"""

    source_file = Path(__file__).parent / "data/shift_ignore/workbook.xlsx"
    output_path = tmp_path / "target.xlsx"
    linking_table_old = Path(__file__).parent / "data/shift_ignore/offsets.csv"
    linking_table_out = tmp_path / "linking_table_out.csv"

    sheet_configs = {
        "page-desc": "skip",
        "page-data": {
            "patient_id_col": "pid",
            "date_columns": ["dob"],
            "text_columns": ["top", "pid"],
            "header_row": 1,
            "skip_rows_after_header": [],
            "shift_ignore": {"dob": ["mssing"]},
        },
    }
    with pytest.raises(ValueError) as info:
        shift_excel_dates_inplace(
            input_file=str(source_file),
            output_file=str(output_path),
            patient_sheet="page-data",
            sheet_configs=sheet_configs,
            min_shift_days=-20,
            max_shift_days=-1,
            seed=14333,
            linking_table_path=str(linking_table_old),
            linking_table_output=str(linking_table_out),
        )
    assert str(info.value) == (
        "patient_id_col='pid' shouldn't be in text_columns of sheet_name='page-data'"
    )


def test_columns_overlap(tmp_path: Path) -> None:
    """check that misconfigurations cause an error"""

    source_file = Path(__file__).parent / "data/shift_ignore/workbook.xlsx"
    output_path = tmp_path / "target.xlsx"
    linking_table_old = Path(__file__).parent / "data/shift_ignore/offsets.csv"
    linking_table_out = tmp_path / "linking_table_out.csv"

    sheet_configs = {
        "page-desc": "skip",
        "page-data": {
            "patient_id_col": "pid",
            "date_columns": ["dob"],
            "text_columns": ["top", "dob"],
            "header_row": 1,
            "skip_rows_after_header": [],
            "shift_ignore": {"dob": ["mssing"]},
        },
    }
    with pytest.raises(ValueError) as info:
        shift_excel_dates_inplace(
            input_file=str(source_file),
            output_file=str(output_path),
            patient_sheet="page-data",
            sheet_configs=sheet_configs,
            min_shift_days=-20,
            max_shift_days=-1,
            seed=14333,
            linking_table_path=str(linking_table_old),
            linking_table_output=str(linking_table_out),
        )
    assert str(info.value) == (
        "sheet_name='page-data' has the some columns in both date and text "
        + "col_names=['dob']"
    )


def test_ptid_column_missing(tmp_path: Path) -> None:
    """check that misconfigurations cause an error"""

    source_file = Path(__file__).parent / "data/shift_ignore/workbook.xlsx"
    output_path = tmp_path / "target.xlsx"
    linking_table_old = Path(__file__).parent / "data/shift_ignore/offsets.csv"
    linking_table_out = tmp_path / "linking_table_out.csv"

    sheet_configs = {
        "page-desc": "skip",
        "page-data": {
            "patient_id_col": "ptid",
            "date_columns": ["dob"],
            "text_columns": ["top", "pid"],
            "header_row": 1,
            "skip_rows_after_header": [],
            "shift_ignore": {"dob": ["mssing"]},
        },
    }
    with pytest.raises(PatientColumnMissing) as info:
        shift_excel_dates_inplace(
            input_file=str(source_file),
            output_file=str(output_path),
            patient_sheet="page-data",
            sheet_configs=sheet_configs,
            min_shift_days=-20,
            max_shift_days=-1,
            seed=14333,
            linking_table_path=str(linking_table_old),
            linking_table_output=str(linking_table_out),
        )
    assert info.value._page_name == "page-data"
    assert info.value._column_name == "ptid"


@pytest.mark.parametrize("shift_ignore", [True, False])
def test_ignore_in_date_columns(shift_ignore: bool, tmp_path: Path) -> None:

    source_file = Path(__file__).parent / "data/shift_ignore/workbook.xlsx"
    output_path = tmp_path / "target.xlsx"
    linking_table_old = Path(__file__).parent / "data/shift_ignore/offsets.csv"
    linking_table_out = tmp_path / "linking_table_out.csv"

    sheet_configs = {
        "page-desc": "skip",
        "page-data": {
            "patient_id_col": "pid",
            "date_columns": ["dob"],
            "text_columns": ["top"],
            "header_row": 1,
            "skip_rows_after_header": [],
        },
    }

    def body() -> None:
        shift_excel_dates_inplace(
            input_file=str(source_file),
            output_file=str(output_path),
            patient_sheet="page-data",
            sheet_configs=sheet_configs,
            min_shift_days=-20,
            max_shift_days=-1,
            seed=14333,
            linking_table_path=str(linking_table_old),
            linking_table_output=str(linking_table_out),
        )

    if not shift_ignore:
        with pytest.raises(ShiftFoundNonDate) as info:
            body()
        assert info.value._page == "page-data"
        assert info.value._row == 5
        assert info.value._col == 1
        assert info.value._col_name == "dob"
        assert info.value._value == "mssing"
        return
    else:
        sheet_configs["page-data"]["shift_ignore"] = {"dob": ["mssing"]}
        body()

    workbook = load_workbook(output_path)

    worksheet = workbook.worksheets[1]

    # first column
    assert worksheet.cell(1, 1).value == "personal id"
    assert worksheet.cell(2, 1).value == "pid"
    assert worksheet.cell(3, 1).value == "nuh71"
    assert worksheet.cell(4, 1).value == "nuh06"
    assert worksheet.cell(5, 1).value == "nuh23"
    assert worksheet.cell(6, 1).value == "nuh67"
    assert worksheet.cell(7, 1).value == "nuh27"

    # check column 3 is all none
    assert not [
        val for val in [worksheet.cell(r + 1, 3).value for r in range(7)] if val
    ]

    # last column
    assert worksheet.cell(1, 4).value == "pizza topping"
    assert worksheet.cell(2, 4).value == "top"
    assert worksheet.cell(3, 4).value == "cheese"
    assert worksheet.cell(4, 4).value == "unknown"
    assert worksheet.cell(5, 4).value == "mushrooms"
    assert worksheet.cell(6, 4).value == "this can't be a date - sorry 2016"
    assert worksheet.cell(7, 4).value == "idk - this can't be a date anymore"

    # the important column to check - the dates
    assert worksheet.cell(1, 2).value == "birthday"
    assert worksheet.cell(2, 2).value == "dob"
    assert str(worksheet.cell(3, 2).value) == "2001-12-17 00:00:00"
    assert str(worksheet.cell(4, 2).value) == "1993-09-20 00:00:00"
    assert worksheet.cell(5, 2).value is None
    assert str(worksheet.cell(6, 2).value) == "mssing"  # change that's under test
    assert str(worksheet.cell(7, 2).value) == "1999-11-30 00:00:00"


@pytest.mark.parametrize("shift_ignore", [True, False])
def test_ignore_in_text_columns(shift_ignore: bool, tmp_path: Path) -> None:

    source_file = Path(__file__).parent / "data/shift_ignore/workbook.xlsx"
    output_path = tmp_path / "target.xlsx"
    linking_table_old = Path(__file__).parent / "data/shift_ignore/offsets.csv"
    linking_table_out = tmp_path / "linking_table_out.csv"

    sheet_configs = {
        "page-desc": {
            "patient_id_col": "pid",
            "date_columns": [],
            "text_columns": ["foo", "glitter"],
            "header_row": 0,
            "skip_rows_after_header": [],
        },
        "page-data": "skip",
    }

    def body(shift_ignore_yaml: None | Path) -> None:
        shift_excel_dates_inplace(
            input_file=str(source_file),
            output_file=str(output_path),
            patient_sheet="page-desc",
            sheet_configs=sheet_configs,
            min_shift_days=-20,
            max_shift_days=-1,
            seed=14333,
            linking_table_path=str(linking_table_old),
            linking_table_output=str(linking_table_out),
            shift_ignore_yaml=shift_ignore_yaml,
        )

    if not shift_ignore:
        with pytest.raises(HiddenDate) as info:
            body(None)
        assert str(info.value._message) == (
            "hidden date in [sheet_name='page-desc', 2, 2 @ glitter] "
            + 'value="can\'t recall the date but on 12/11/2001 they had an itchy tummy"'
            + " // found=datetime.datetime(2001, 12, 11, 0, 0)"
        )
        return

    else:
        # add the exception

        sheet_configs["page-desc"]["shift_ignore"] = {
            "glitter": [
                "can't recall the date but on 12/11/2001 they had an itchy tummy",
            ],
            "foo": [
                "agreed on 21st jul 2017",
            ],
        }

        # run the shift
        body(source_file)

        workbook = load_workbook(output_path)
        page = workbook["page-desc"]

        obtained = [
            [str(page.cell(row + 1, col + 1).value) for col in range(page.max_column)]
            for row in range(page.max_row)
        ]

        expected = [
            ["foo", "pid", "glitter"],
            ["bar", "nuh71", "road"],
            [
                "shoe",
                "nuh06",
                "can't recall the date but on 12/11/2001 they had an itchy tummy",
            ],
            ["agreed on 21st jul 2017", "nuh23", "grip"],
            ["farm", "nuh67", "grim"],
            ["cake", "nuh27", "cheese"],
            ["glip", "nuh65", "jkl"],
        ]

        assert expected == obtained


def test_move_to_file(tmp_path: Path) -> None:

    source_file = Path(__file__).parent / "data/shift_ignore/workbook.xlsx"
    output_path = tmp_path / "target.xlsx"
    linking_table_old = Path(__file__).parent / "data/shift_ignore/offsets.csv"
    linking_table_out = tmp_path / "linking_table_out.csv"

    sheet_configs = {
        "page-desc": {
            "patient_id_col": "pid",
            "date_columns": [],
            "text_columns": ["foo", "glitter"],
            "header_row": 0,
            "skip_rows_after_header": [],
        },
        "page-data": "skip",
    }

    # add the old config we want to remove
    sheet_configs["page-desc"]["shift_ignore"] = {
        "glitter": [
            "can't recall the date but on 12/11/2001 they had an itchy tummy",
        ],
        "foo": [
            "agreed on 21st jul 2017",
        ],
    }

    with pytest.raises(RuntimeError) as error:
        shift_excel_dates_inplace(
            input_file=str(source_file),
            output_file=str(output_path),
            patient_sheet="page-desc",
            sheet_configs=sheet_configs,
            min_shift_days=-20,
            max_shift_days=-1,
            seed=14333,
            linking_table_path=str(linking_table_old),
            linking_table_output=str(linking_table_out),
        )
        pytest.fail("execution should throw an exception before now")

    assert str(error.value) == "move shift_ignore from sheet_configs to a file"
